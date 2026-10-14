"""Bounded JSON/XLSX import; all outputs use the desktop's canonical validator."""
from io import BytesIO
import json
from pathlib import Path
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from proctoring.exam.quiz import normalize_question_set

MAX_QUESTIONS = 200
MAX_ROWS = 5001
XLSX_COLUMNS = ("question_id", "type", "prompt", "option_id", "option_text", "is_correct")


def unique_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"Duplicate JSON field: {key}")
        result[key] = value
    return result


def normalized(data: object) -> dict:
    if not isinstance(data, dict) or not isinstance(data.get("questions"), list):
        raise ValueError("Import must contain a title and questions array")
    if not 1 <= len(data["questions"]) <= MAX_QUESTIONS:
        raise ValueError(f"Import must contain 1–{MAX_QUESTIONS} questions")
    if len(json.dumps(data, ensure_ascii=False)) > 2 * 1024 * 1024:
        raise ValueError("Question set exceeds the 2 MiB content limit")
    return normalize_question_set(data)


def parse_import(filename: str, contents: bytes, max_bytes: int = 2 * 1024 * 1024) -> dict:
    if not contents or len(contents) > max_bytes:
        raise ValueError("Import must be nonempty and at most 2 MiB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        try:
            return normalized(json.loads(contents.decode("utf-8-sig"), object_pairs_hook=unique_object,
                                         parse_constant=lambda value: (_ for _ in ()).throw(ValueError("Non-finite JSON numbers are not allowed"))))
        except (UnicodeError, json.JSONDecodeError, RecursionError) as error:
            raise ValueError("Invalid UTF-8 JSON question set") from error
    if suffix != ".xlsx":
        raise ValueError("Choose a .json or .xlsx question set")
    try:
        with ZipFile(BytesIO(contents)) as archive:
            members = archive.infolist()
            if len(members) > 100 or sum(item.file_size for item in members) > 16 * 1024 * 1024:
                raise ValueError("XLSX archive expands beyond the supported import limit")
            for item in members:
                lower = item.filename.lower()
                if (item.flag_bits & 1 or "vbaproject" in lower or "externallinks/" in lower
                        or lower.endswith(".bin") or "../" in lower
                        or item.file_size > max(1024 * 1024, item.compress_size * 200)):
                    raise ValueError("Encrypted, macro, external-link or excessive-compression XLSX files are not supported")
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=False, keep_links=False)
        try:
            if workbook.sheetnames != ["Questions"]:
                raise ValueError("XLSX must contain exactly one sheet named Questions")
            sheet = workbook["Questions"]
            if (sheet.max_row is not None and sheet.max_row > MAX_ROWS or
                    sheet.max_column is not None and sheet.max_column > len(XLSX_COLUMNS)):
                raise ValueError("XLSX contains too many rows or columns")
            rows = sheet.iter_rows()
            header = tuple(cell.value for cell in next(rows))
            if header != XLSX_COLUMNS:
                raise ValueError("XLSX header must be: " + ", ".join(XLSX_COLUMNS))
            questions: dict[str, dict] = {}
            for row_number, cells in enumerate(rows, 2):
                if row_number > MAX_ROWS:
                    raise ValueError("XLSX contains too many rows")
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError(f"Row {row_number}: formulas are not allowed; use literal values")
                values = [cell.value for cell in cells]
                if all(value is None for value in values):
                    continue
                qid, kind, prompt, oid, text, correct = values
                if any(not isinstance(value, str) or not value.strip() for value in (qid, kind, prompt, oid, text)):
                    raise ValueError(f"Row {row_number}: IDs, type, prompt and option text must be nonempty text")
                if not (type(correct) is bool or type(correct) is int and correct in (0, 1)):
                    raise ValueError(f"Row {row_number}: is_correct must be TRUE/FALSE or 1/0")
                question = questions.setdefault(qid, {"id": qid, "type": kind, "prompt": prompt,
                                                       "options": [], "correct_option_ids": []})
                if question["type"] != kind or question["prompt"] != prompt:
                    raise ValueError(f"Row {row_number}: repeated question ID has a different type or prompt")
                question["options"].append({"id": oid, "text": text})
                if correct:
                    question["correct_option_ids"].append(oid)
            return normalized({"title": Path(filename).stem, "questions": list(questions.values())})
        finally:
            workbook.close()
    except (BadZipFile, KeyError, OSError, StopIteration, ParseError) as error:
        raise ValueError("Invalid XLSX workbook") from error
