"""Versioned reviewed PDFs and typed XLSX exports, never replacements for originals."""
from __future__ import annotations

from hashlib import sha256
from html import escape
from io import BytesIO
import json
import os
from pathlib import Path
import threading
from uuid import uuid4

from .service import Problem, iso

_FONT_LOCK = threading.Lock()
_LABELS = {'not_reviewed': 'Not reviewed', 'confirmed': 'Observation confirmed',
           'false_positive': 'False positive', 'insufficient_evidence': 'Insufficient evidence',
           'none_confirmed': 'No violations confirmed', 'violation_confirmed': 'Violation confirmed by teacher'}


def _font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    with _FONT_LOCK:
        if 'ProctorReview' not in pdfmetrics.getRegisteredFontNames():
            paths = [Path(os.environ.get('PROCTOR_REPORT_FONT', '')), Path('C:/Windows/Fonts/arial.ttf'),
                     Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]
            path = next((p for p in paths if p.is_file()), None)
            if path is None:
                raise RuntimeError('A Unicode TrueType font is required; configure PROCTOR_REPORT_FONT')
            pdfmetrics.registerFont(TTFont('ProctorReview', str(path)))
    return 'ProctorReview'


def build_reviewed_pdf(snapshot, image_lookup):
    """Called by a background task, using a frozen conclusion version."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, Image

    font, output = _font(), BytesIO()
    body = ParagraphStyle('Body', fontName=font, fontSize=8.5, leading=12, spaceAfter=5, wordWrap='CJK')
    title = ParagraphStyle('Title', parent=body, fontSize=20, leading=25, spaceAfter=12, textColor=colors.HexColor('#355b38'))
    heading = ParagraphStyle('Heading', parent=body, fontSize=12, leading=16, spaceBefore=13, spaceAfter=8, keepWithNext=True)
    small = ParagraphStyle('Small', parent=body, fontSize=7.5, leading=10)
    def paragraph(value, style=body):
        return Paragraph(escape(str(value if value is not None else '—')).replace('\n', '<br/>'), style)
    def table(rows, widths):
        result = LongTable(rows, colWidths=widths, repeatRows=1, hAlign='LEFT', splitInRow=1)
        result.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e9eee7')),
            ('VALIGN',(0,0),(-1,-1),'TOP'), ('BOX',(0,0),(-1,-1),.4,colors.HexColor('#d8ddd4')),
            ('INNERGRID',(0,0),(-1,-1),.25,colors.HexColor('#d8ddd4')),
            ('LEFTPADDING',(0,0),(-1,-1),6), ('RIGHTPADDING',(0,0),(-1,-1),6),
            ('TOPPADDING',(0,0),(-1,-1),6), ('BOTTOMPADDING',(0,0),(-1,-1),6)]))
        return result
    attempt, exam, sync, conclusion = (snapshot[k] for k in ('attempt','exam','sync','conclusion'))
    timing = sync.get('timing') or {}
    flow = [paragraph('Reviewed exam report', title), paragraph(f"Version {conclusion['version']} • {exam['title']}"),
            paragraph('A confirmed observation does not itself establish an academic violation.'),
            paragraph(f"Student: {attempt['full_name']}\nGroup: {attempt['group']}\nInstitutional ID: {attempt['student_id']}\nAttempt: {attempt['sequence']} ({attempt['id']})"),
            paragraph(f"Started: {timing.get('started_at') or attempt.get('started_at')}\nEnded: {timing.get('ended_at') or attempt.get('finished_at')}\nActive exam time: {timing.get('elapsed_seconds', 'Unavailable')} seconds\nDuration: {exam['duration_seconds']} seconds"),
            paragraph('Results and integrity', heading)]
    verified = sync.get('server_result') or {}
    if exam['exam_type'] == 'external':
        flow.append(paragraph('External exam. No LMS grade is available; local completion does not submit external answers.'))
    else:
        flow.append(paragraph(f"Local/client-reported: {attempt.get('score')} / {attempt.get('total')}\nServer-verified answers: {verified.get('score')} / {verified.get('total')}\nConsistent with local result: {verified.get('consistent')}\nRecalculation does not establish that the student computer was tamper-proof."))
    flow.extend([paragraph('Teacher conclusion', heading), paragraph(_LABELS[conclusion['decision']]),
                 paragraph(conclusion['comment']),
                 paragraph(f"Unreviewed behavioral observations at finalization: {conclusion['unreviewed_count']}\nReview author: {conclusion.get('teacher_name', conclusion['teacher_id'])}\nFinalized: {conclusion['created_at']}"),
                 paragraph('Event timeline', heading)])
    rows = [[paragraph(x, small) for x in ('Time (UTC)','Event / category','Duration (s)','Teacher decision / comment','Evidence')]]
    for event in snapshot['events']:
        review = event['review']
        evidence = []
        for file in event['files']:
            path = image_lookup(file['id'])
            image = Image(str(path))
            factor = min(80/image.imageWidth, 56/image.imageHeight)
            image.drawWidth, image.drawHeight = image.imageWidth*factor, image.imageHeight*factor
            evidence.extend([image, paragraph(file['id'], small)])
        from datetime import datetime, timezone
        try:
            stamp = datetime.fromisoformat(event['timestamp'].replace('Z','+00:00')).astimezone(timezone.utc).strftime('%H:%M:%S')
        except (AttributeError,ValueError):
            stamp = event['timestamp']
        rows.append([paragraph(stamp, small), paragraph(f"{event['event_type']}\n{event['category']}", small),
                     paragraph(event['duration_seconds'], small),
                     paragraph(f"{_LABELS[review['decision']]}\n{review['comment']}\n{review.get('teacher_name') or ''}\n{review.get('created_at') or ''}", small),
                     evidence or paragraph('No capture', small)])
    flow.append(table(rows, [70,112,50,164,100]) if len(rows)>1 else paragraph('No automatic observations or technical records received.'))
    flow.append(paragraph('Manual teacher observations', heading))
    if not snapshot['manual_observations']:
        flow.append(paragraph('None.'))
    for item in snapshot['manual_observations']:
        flow.append(paragraph(f"{item['category']} — {'Cancelled' if item['cancelled'] else 'Active'}; version {item['version']}"))
        flow.append(paragraph(item['description']))
        flow.append(paragraph(f"Author: {item['author_name']}\nCreated: {item['created_at']} ({item['phase']})\nLast change: {item['updated_at']}\nManual observation; no screenshot was created.", small))
    reference = next((f for f in sync.get('files', []) if f.get('kind') == 'reference' and f.get('complete')), None)
    if reference:
        image = Image(str(image_lookup(reference['id'])))
        factor = min(140/image.imageWidth, 105/image.imageHeight)
        image.drawWidth, image.drawHeight = image.imageWidth*factor, image.imageHeight*factor
        flow.extend([paragraph('Available identity reference', heading), image, paragraph(reference['id'], small)])
    flow.extend([paragraph('Evidence receipt', heading), paragraph(json.dumps(sync.get('receipt'), ensure_ascii=False), small),
                 paragraph('This version is an immutable review snapshot. Later corrections may require reconsideration; consult the teacher dashboard for the current version.', small)])
    def footer(canvas, document):
        canvas.saveState(); canvas.setFont(font, 8)
        canvas.drawString(50, 25, f"Attempt {attempt['sequence']} · review v{conclusion['version']}")
        canvas.drawRightString(A4[0]-50, 25, str(document.page)); canvas.restoreState()
    SimpleDocTemplate(output, pagesize=A4, leftMargin=50, rightMargin=49, topMargin=42, bottomMargin=45,
                      title='Reviewed exam report', author=conclusion.get('teacher_name', 'Teacher')).build(flow, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


class ReviewReports:
    def __init__(self, review):
        self.review, self.service, self.sync = review, review.service, review.sync
        self._lock = threading.RLock()

    def _path(self, attempt_id, version):
        if not isinstance(attempt_id, str) or len(attempt_id) != 32 or any(c not in '0123456789abcdef' for c in attempt_id):
            raise Problem('not_found', 'Report not found', 404)
        root = self.service.settings.evidence_root.resolve()
        path = root/attempt_id/'reviewed'/f'review-{version}.pdf'
        if path.is_symlink() or any(parent.is_symlink() for parent in (path.parent,path.parent.parent)) or not path.resolve().is_relative_to(root):
            raise Problem('invalid_path', 'Report storage path is unavailable', 409)
        return path

    def generate(self, owner, attempt_id, version):
        """Atomic version output. Recheck deletion under the same DB write lock."""
        with self._lock:
            try:
                with self.service.database.connect() as connection:
                    self.review._owned(connection, owner, attempt_id)
                    row = connection.execute('SELECT * FROM attempt_conclusions WHERE attempt_id=? AND version=?', (attempt_id,version)).fetchone()
                    if not row:
                        return
                    if row['report_status'] == 'ready':
                        return
                    snapshot = json.loads(row['review_snapshot'])
                raw = build_reviewed_pdf(snapshot, lambda item: self.sync.evidence_file(owner,attempt_id,item)[0])
                path = self._path(attempt_id,version)
                with self.service.database.connect(write=True) as connection:
                    self.review._owned(connection,owner,attempt_id)
                    # A second worker/process may have completed this immutable
                    # version while the PDF was being rendered outside the lock.
                    current = connection.execute('SELECT report_status FROM attempt_conclusions WHERE attempt_id=? AND version=?', (attempt_id,version)).fetchone()
                    if not current or current['report_status'] in ('ready','deleted'):
                        return
                    self.sync.require_complete(connection,attempt_id)
                    path.parent.mkdir(parents=True,exist_ok=True)
                    temporary = path.with_name('.'+path.name+'.'+uuid4().hex+'.partial')
                    try:
                        with temporary.open('xb') as handle:
                            handle.write(raw); handle.flush(); os.fsync(handle.fileno())
                        os.replace(temporary,path)
                    finally:
                        temporary.unlink(missing_ok=True)
                    connection.execute("UPDATE attempt_conclusions SET report_status='ready',report_sha256=?,report_size=?,report_error=NULL WHERE attempt_id=? AND version=?",
                                       (sha256(raw).hexdigest(),len(raw),attempt_id,version))
            except Exception:
                # Do not expose paths, raw student data or dependency exception text.
                with self.service.database.connect(write=True) as connection:
                    connection.execute("UPDATE attempt_conclusions SET report_status='failed',report_error='Reviewed PDF creation failed; retry after checking the server report dependencies.' WHERE attempt_id=? AND version=? AND report_status NOT IN ('ready','deleted')", (attempt_id,version))

    def reviewed_file(self, owner, attempt_id, version):
        with self.service.database.connect() as connection:
            self.review._owned(connection,owner,attempt_id)
            row = connection.execute('SELECT * FROM attempt_conclusions WHERE attempt_id=? AND version=?', (attempt_id,version)).fetchone()
            if not row or row['report_status'] != 'ready':
                raise Problem('report_not_ready','This reviewed PDF is not available yet',409)
            path = self._path(attempt_id,version)
            if not path.is_file() or path.stat().st_size != row['report_size'] or sha256(path.read_bytes()).hexdigest() != row['report_sha256']:
                raise Problem('report_corrupt','Reviewed PDF failed its integrity check',409)
            return path

    def original_file(self, owner, attempt_id):
        detail = self.sync.teacher_detail(owner,attempt_id)
        file = next((f for f in detail.get('files',[]) if f['kind']=='original_pdf' and f['complete']),None)
        if not file:
            raise Problem('report_not_ready','The original local PDF has not been fully received',409)
        return self.sync.evidence_file(owner,attempt_id,file['id'])[0]

    def xlsx(self, owner, exam_id):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from datetime import datetime, timezone
        with self.service.database.connect() as connection:
            exam = self.service.owned_exam(connection,owner,exam_id)
            attempts = [dict(r) for r in connection.execute('SELECT * FROM attempts WHERE exam_id=? ORDER BY admitted_at,id',(exam_id,))]
        book = Workbook(); sheet = book.active; sheet.title='Attempts'
        headers = ['Student full name','Group','Institutional Student ID','Attempt number','Status','Start (UTC)','End (UTC)',
                   'Client score','Server-verified score','Total questions','Server percentage','Behavioral events','Technical events',
                   'Review status','Final conclusion','Result integrity']
        sheet.append(headers)
        for attempt in attempts:
            try:
                detail = self.review.detail(owner,attempt['id'])
            except Problem as error:
                if error.code != 'materials_deleted':
                    raise
                detail = None
            conclusion = detail['conclusion'] if detail else None
            sync = detail['sync'] if detail else {}
            result = sync.get('server_result') or {}
            verified = result.get('score',result.get('verified_score')) if exam['exam_type']=='native' else None
            total = result.get('total',attempt['total']) if exam['exam_type']=='native' else None
            started = (sync.get('timing') or {}).get('started_at')
            ended = (sync.get('timing') or {}).get('ended_at')
            def date(value, fallback):
                if value:
                    return datetime.fromisoformat(value.replace('Z','+00:00')).astimezone(timezone.utc).replace(tzinfo=None)
                return datetime.fromtimestamp(fallback,timezone.utc).replace(tzinfo=None) if fallback is not None else None
            row = [attempt['full_name'],attempt['student_group'],attempt['student_id'],attempt['sequence'],
                   sync.get('status','materials_deleted' if detail is None else detail['attempt']['status']),
                   date(started,attempt['started_at']),date(ended,attempt['finished_at']),
                   attempt['score'] if exam['exam_type']=='native' else None,verified,total,
                   verified/total if type(verified) in (int,float) and total else None,
                   sum(e['category']=='behavioral' for e in detail['events']) if detail else None,
                   sum(e['category']=='technical' for e in detail['events']) if detail else None,
                   ('needs_reconsideration' if conclusion['needs_reconsideration'] else 'finalized') if conclusion else 'not_finalized',
                   _LABELS.get(conclusion['decision']) if conclusion else None,
                   'mismatch' if result.get('consistent') is False else result.get('source','not_verified')]
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value,str):
                    cell.data_type='s'  # Prevent formulas from teacher/student-supplied strings.
                cell.alignment=Alignment(vertical='top',wrap_text=True)
            for column in (6,7):
                sheet.cell(sheet.max_row,column).number_format='yyyy-mm-dd hh:mm:ss'
            sheet.cell(sheet.max_row,11).number_format='0.0%'
        for cell in sheet[1]:
            cell.font=Font(bold=True,color='355B38'); cell.fill=PatternFill('solid',fgColor='E9EEE7')
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        widths=[28,18,24,16,24,23,23,15,21,18,19,20,18,25,38,24]
        for index,width in enumerate(widths,1):
            sheet.column_dimensions[sheet.cell(1,index).column_letter].width=width
        sheet.row_dimensions[1].height=34
        sheet.freeze_panes='D2'; sheet.auto_filter.ref=sheet.dimensions
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation='landscape'; sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth=1; sheet.page_setup.fitToHeight=0
        sheet.print_title_rows='1:1'
        output=BytesIO(); book.save(output); return output.getvalue()
