from concurrent.futures import ThreadPoolExecutor
from hashlib import sha256
from io import BytesIO
import json

import pytest

from proctoring_server.review import ReviewService, event_groups, event_projection
from proctoring_server.review_reports import ReviewReports
from proctoring_server.service import Problem
from review_fixtures import build_review_bundle


@pytest.fixture
def bundle(service,owner,exam_data):
    return build_review_bundle(service,owner,exam_data)


@pytest.fixture
def review(service,bundle):
    return ReviewService(service,bundle['sync'])


def decision(event_id, expected=0, value='confirmed', comment='Observed phone; no academic conclusion.'):
    return {'event_id':event_id,'expected_version':expected,'decision':value,'comment':comment}


def conclusion(expected=0,confirm=True,value='none_confirmed'):
    return {'expected_version':expected,'decision':value,'comment':'Reviewed available evidence.','confirm_unreviewed':confirm}


def test_projection_groups_without_destroying_raw_versions(review,bundle,owner):
    detail=review.detail(owner,bundle['attempt_id'])
    assert len(detail['events'])==3
    assert detail['events'][0]['duration_seconds']==2
    assert detail['groups'][0]['event_ids']==['event-1','event-2']
    assert detail['events'][2]['category']=='technical'
    assert detail['unreviewed_count']==2
    assert detail['events'][0]['files']==[{'id':'snapshot','kind':'snapshot'}]
    assert len(bundle['sync'].records(owner,bundle['attempt_id']))==4


def test_actual_lifecycle_outage_duration_and_point_events():
    records = [{'sequence':i,'record':record} for i,record in enumerate([
        {'event_type':'monitoring_failed','timestamp':'2026-10-14T10:00:00+00:00'},
        {'event_type':'monitoring_recovered','timestamp':'2026-10-14T10:00:20+00:00','duration':20.125},
        {'event_type':'monitoring_interruption_ended','timestamp':'2026-10-14T10:01:00+00:00','duration':4.5},
        {'event_type':'phone_visible','timestamp':'2026-10-14T10:02:00+00:00','duration_seconds':3.75},
        {'event_type':'session_ended','timestamp':'2026-10-14T10:03:00+00:00','duration_seconds':False},
    ])]
    assert [e['duration_seconds'] for e in event_projection(records, [])] == [None,20.125,4.5,3.75,None]


def test_presentation_groups_fit_atomic_bulk_review_limit():
    events = [{'event_id':f'event-{i}','event_type':'phone_visible','category':'behavioral',
               'timestamp':'2026-10-14T10:00:00+00:00','duration_seconds':1.} for i in range(101)]
    groups = event_groups(events)
    assert [len(g['event_ids']) for g in groups] == [100,1]
    assert [key for group in groups for key in group['event_ids']] == [e['event_id'] for e in events]


def test_bulk_review_individual_exception_and_audit(review,bundle,owner,service):
    attempt=bundle['attempt_id']
    review.review(owner,attempt,{'items':[decision('event-1'),decision('event-2')]})
    detail=review.review(owner,attempt,{'items':[decision('event-2',1,'false_positive','Reflection, not a phone.')]})
    assert [e['review']['decision'] for e in detail['events'][:2]]==['confirmed','false_positive']
    assert len(detail['events'][1]['history'])==2
    assert detail['conclusion'] is None
    with service.database.connect() as connection:
        rows=connection.execute("SELECT details FROM audit_log WHERE action='event_reviewed' AND resource_id=?",(attempt,)).fetchall()
    changed=json.loads(rows[-1][0]);assert changed['old']['decision']=='confirmed' and changed['new']['decision']=='false_positive'


def test_bulk_conflict_is_atomic(review,bundle,owner):
    attempt=bundle['attempt_id']
    review.review(owner,attempt,{'items':[decision('event-2')]})
    with pytest.raises(Problem,match='changed'):
        review.review(owner,attempt,{'items':[decision('event-1'),decision('event-2')]})
    assert review.detail(owner,attempt)['events'][0]['review']['version']==0


@pytest.mark.parametrize('bad',[{'event_id':[],'expected_version':0,'decision':'confirmed'},decision('event-1',value=[]),decision('missing')])
def test_bad_reviews_rejected(review,bundle,owner,bad):
    with pytest.raises(Problem):review.review(owner,bundle['attempt_id'],{'items':[bad]})


def test_all_teacher_material_and_review_access_is_isolated(review,bundle,owner,other):
    attempt=bundle['attempt_id'];reports=ReviewReports(review)
    for action in [lambda:review.detail(other,attempt),lambda:review.review(other,attempt,{'items':[decision('event-1')]}),
                   lambda:review.manual(other,attempt,{'category':'note','description':'Forbidden'}),
                   lambda:review.conclude(other,attempt,conclusion()),lambda:reports.original_file(other,attempt),
                   lambda:reports.xlsx(other,bundle['exam_id'])]:
        with pytest.raises(Problem) as failure:action()
        assert failure.value.status==404


def test_manual_history_cancellation_invalidates_final_conclusion(review,bundle,owner,clock):
    attempt=bundle['attempt_id']
    detail=review.manual(owner,attempt,{'category':'Teacher note','description':'Initial description'})
    item=detail['manual_observations'][0]
    assert item['phase']=='later_review'
    review.conclude(owner,attempt,conclusion())
    clock.advance(2)
    detail=review.manual(owner,attempt,{'expected_version':1,'category':'Teacher note','description':'Corrected description','cancelled':True},item['id'])
    item=detail['manual_observations'][0]
    assert item['cancelled'] and len(item['history'])==2 and item['created_at']!=item['updated_at']
    assert item['history'][0]['description']=='Initial description'
    assert detail['conclusion']['needs_reconsideration']
    assert detail['events'][0]['files']==[{'id':'snapshot','kind':'snapshot'}] # no fake manual screenshot


def test_manual_creation_during_recorded_exam(service,owner,exam_data,clock):
    # Build a real admitted/confirmed start, then inspect before a final manifest.
    import secrets
    from uuid import uuid4
    from proctoring_server.sync import SyncService
    from proctoring_server.service import iso
    exam=service.create_exam(owner,exam_data);cred=secrets.token_urlsafe(32);device=secrets.token_urlsafe(32)
    admission=service.admit({'exam_code':exam['code'],'full_name':'Student','group':'A','student_id':'during',
        'admission_id':uuid4().hex,'device_id':uuid4().hex,'device_secret':device,'credential':cred})
    sync=SyncService(service); start={'start_key':uuid4().hex,'session_id':'session-'+uuid4().hex}
    sync.authorize_start(cred,device,start);sync.confirm_start(cred,device,{**start,'started_at':iso(clock())})
    detail=ReviewService(service,sync).manual(owner,admission['attempt_id'],{'category':'In-room observation','description':'Teacher observed a raised hand.'})
    assert detail['manual_observations'][0]['phase']=='during_exam'
    assert not detail['can_finalize']


def test_ended_client_flag_without_receipt_cannot_finalize(service,owner,exam_data):
    bundle=build_review_bundle(service,owner,exam_data,complete=False)
    review=ReviewService(service,bundle['sync']);attempt=bundle['attempt_id']
    assert review.detail(owner,attempt)['attempt']['finished_at'] is not None
    assert not review.detail(owner,attempt)['can_finalize']
    with pytest.raises(Problem):review.conclude(owner,attempt,conclusion())


def test_unreviewed_requires_explicit_confirmation_and_keeps_score(review,bundle,owner,service):
    attempt=bundle['attempt_id']
    with pytest.raises(Problem) as failure:review.conclude(owner,attempt,conclusion(confirm=False))
    assert failure.value.code=='unreviewed_confirmation_required'
    review.conclude(owner,attempt,conclusion(value='violation_confirmed'))
    detail=review.detail(owner,attempt)
    assert detail['conclusion']['unreviewed_count']==2
    assert detail['attempt']['score']==1
    assert detail['sync']['server_result']['score']==2
    assert detail['sync']['server_result']['consistent'] is False


def test_pdf_versions_unicode_original_preserved_and_typed_xlsx(review,bundle,owner):
    from pypdf import PdfReader
    from openpyxl import load_workbook
    attempt=bundle['attempt_id'];reports=ReviewReports(review)
    original=reports.original_file(owner,attempt).read_bytes()
    assert original==bundle['raw']['original_pdf']
    first=review.conclude(owner,attempt,conclusion())
    reports.generate(owner,attempt,first)
    pdf1=reports.reviewed_file(owner,attempt,first).read_bytes()
    text='\n'.join(page.extract_text() for page in PdfReader(BytesIO(pdf1)).pages)
    assert 'Әлия Иванова' in text and 'Unreviewed behavioral observations at finalization: 2' in text
    assert 'monitoring_failed' in text and 'No violations confirmed' in text
    review.review(owner,attempt,{'items':[decision('event-1')]})
    assert review.detail(owner,attempt)['conclusion']['needs_reconsideration']
    second=review.conclude(owner,attempt,conclusion(1,value='insufficient_evidence'))
    reports.generate(owner,attempt,second)
    assert reports.reviewed_file(owner,attempt,first).read_bytes()==pdf1
    assert reports.reviewed_file(owner,attempt,second).read_bytes()!=pdf1
    assert reports.original_file(owner,attempt).read_bytes()==original
    workbook=load_workbook(BytesIO(reports.xlsx(owner,bundle['exam_id'])))
    sheet=workbook.active
    assert sheet['B2'].value=='=1+1' and sheet['B2'].data_type=='s'
    assert sheet['C2'].value=='00123'
    assert sheet['I2'].value==2 and sheet['K2'].value==1
    assert sheet['F2'].data_type=='d'


def test_report_failure_is_visible_and_retry_recovers(review,bundle,owner,monkeypatch):
    import proctoring_server.review_reports as module
    attempt=bundle['attempt_id'];reports=ReviewReports(review);version=review.conclude(owner,attempt,conclusion())
    original=module.build_reviewed_pdf
    monkeypatch.setattr(module,'build_reviewed_pdf',lambda *args:(_ for _ in ()).throw(OSError('test only')))
    reports.generate(owner,attempt,version)
    assert review.detail(owner,attempt)['conclusion']['report_status']=='failed'
    with pytest.raises(Problem):reports.reviewed_file(owner,attempt,version)
    monkeypatch.setattr(module,'build_reviewed_pdf',original)
    reports.generate(owner,attempt,version)
    assert reports.reviewed_file(owner,attempt,version).stat().st_size>1000


def test_competing_report_workers_cannot_replace_ready_version(review,bundle,owner,monkeypatch):
    import proctoring_server.review_reports as module
    attempt=bundle['attempt_id']; first=ReviewReports(review); second=ReviewReports(review)
    version=review.conclude(owner,attempt,conclusion())
    original=module.build_reviewed_pdf
    winning=[]
    def competing_render(snapshot, lookup):
        # First worker is rendering. Another instance commits the same version
        # before the first worker reaches its final serialized write transaction.
        monkeypatch.setattr(module,'build_reviewed_pdf',original)
        second.generate(owner,attempt,version)
        winning.append(second.reviewed_file(owner,attempt,version).read_bytes())
        return b'%PDF-intentionally-different-losing-worker'
    monkeypatch.setattr(module,'build_reviewed_pdf',competing_render)
    first.generate(owner,attempt,version)
    assert first.reviewed_file(owner,attempt,version).read_bytes()==winning[0]


def test_deleted_material_prevents_late_report_resurrection(review,bundle,owner):
    attempt=bundle['attempt_id'];reports=ReviewReports(review);version=review.conclude(owner,attempt,conclusion())
    bundle['sync'].delete_materials(owner,attempt,{'confirm':True,'reason':'Synthetic deletion test'})
    reports.generate(owner,attempt,version)
    with pytest.raises(Problem):review.detail(owner,attempt)
    assert not (review.service.settings.evidence_root/attempt/'reviewed').exists()


def test_concurrent_review_edits_do_not_lose_history(review,bundle,owner):
    def save(comment):
        try:
            review.review(owner,bundle['attempt_id'],{'items':[decision('event-1',comment=comment)]});return 'ok'
        except Problem as error:return error.code
    with ThreadPoolExecutor(2) as pool:results=list(pool.map(save,['first','second']))
    assert sorted(results)==['ok','review_conflict']


def test_corrupt_evidence_blocks_new_final_conclusion(review,bundle,owner):
    path=bundle['sync'].file_path(owner,bundle['attempt_id'],'snapshot')
    path.write_bytes(b'corrupted after initial receipt')
    assert not review.detail(owner,bundle['attempt_id'])['can_finalize']
    with pytest.raises(Problem):review.conclude(owner,bundle['attempt_id'],conclusion())


def test_actual_review_routes_require_teacher_cookie_csrf_and_preserve_downloads(client,bundle,owner):
    origin={'Origin':'http://127.0.0.1:5174'}
    auth=client.post('/api/v1/auth/login',headers=origin,json={'email':'teacher@example.com','password':'test password 12345'})
    headers={**origin,'X-CSRF-Token':auth.json()['csrf_token']}
    base='/api/v1/attempts/'+bundle['attempt_id']
    assert client.get(base+'/review').status_code==200
    assert client.post(base+'/reviews',json={'items':[decision('event-1')]}).status_code==403
    assert client.post(base+'/reviews',headers=headers,json={'items':[decision('event-1')]}).status_code==200
    response=client.post(base+'/conclusions',headers=headers,json=conclusion())
    assert response.status_code==202,response.text
    original=client.get(base+'/reports/original')
    assert original.content==bundle['raw']['original_pdf']
    reviewed=client.get(base+'/reports/reviewed/1')
    assert reviewed.status_code==200 and reviewed.content.startswith(b'%PDF-')
    assert reviewed.headers['cache-control']=='no-store'
    exported=client.get('/api/v1/exams/'+bundle['exam_id']+'/export.xlsx')
    assert exported.status_code==200 and exported.content.startswith(b'PK')
    client.cookies.clear()
    student={'Authorization':'Bearer '+bundle['credential'],'X-Device-Secret':bundle['device_secret']}
    assert client.get(base+'/review',headers=student).status_code==401
    assert client.get(base+'/reports/reviewed/1',headers=student).status_code==401


def test_external_xlsx_never_invents_grades(service,owner,exam_data):
    from openpyxl import load_workbook
    from proctoring_server.sync import SyncService
    data={**exam_data,'exam_type':'external','external_url':'https://example.edu/exam'}
    exam=service.create_exam(owner,data)
    admission=service.admit({'exam_code':exam['code'],'full_name':'External Student','group':'A','student_id':'ext'})
    service.start(admission['credential']);service.finish(admission['credential'],{'score':None,'total':None,'end_reason':'completed'})
    reports=ReviewReports(ReviewService(service,SyncService(service)))
    sheet=load_workbook(BytesIO(reports.xlsx(owner,exam['id']))).active
    assert all(sheet.cell(2,col).value is None for col in (8,9,10,11))
